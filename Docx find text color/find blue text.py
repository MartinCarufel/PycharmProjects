from docx import Document
from docx.shared import RGBColor

doc = Document("text.docx")


def get_paragraph(doc):

    for para in doc.paragraphs:
        for run in para.runs:
            if run.font.color.rgb == RGBColor(0, 176, 240):
                print(run.text)


def get_table(doc):
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.font.color.rgb == RGBColor(0, 176, 240):

                            print(run.text)

get_paragraph(doc)